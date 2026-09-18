import asyncio
import os
import json
import uuid
from typing import List
from sqlalchemy.orm import Session
from fastapi import UploadFile, HTTPException
from ..models.document import Document, DocumentChunk
from ..core.config import settings
from ..core.database import SessionLocal
from ..rag.embeddings import embed_texts
from . import plan_service, web_crawl
# The SSRF guard, robots/sitemap parsing, and page-discovery crawler used to
# be defined here; extracted to web_crawl.py in Stage 2 of apps/agents/
# seo-copywriter/CLAUDE.md so the SEO Audit & Optimization agent's own
# crawler reuses this exact hardened fetch path instead of a second
# implementation. Re-exported under their old names so every existing call
# site and test in this module keeps working unchanged.
from .web_crawl import (
    assert_public_url as _assert_public_url,
    discover_website_pages,
    discover_sitemap_urls as _discover_sitemap_urls,
    discover_by_crawling as _discover_by_crawling,
    fetch_sitemap_xml as _fetch_sitemap_xml,
    get_robot_parser as _get_robot_parser,
    looks_like_page as _looks_like_page,
    site_root as _site_root,
    CRAWL_USER_AGENT,
    MAX_CRAWL_PAGES,
    MAX_FETCH_BYTES as MAX_URL_FETCH_BYTES,
)


ALLOWED_TYPES = {"pdf", "docx", "txt", "csv", "xlsx", "url"}


def _extract_text(path: str, file_type: str) -> str:
    if file_type == "txt" or file_type == "csv":
        with open(path, encoding="utf-8", errors="ignore") as f:
            return f.read()
    if file_type == "pdf":
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(path)
            return "\n".join(p.extract_text() or "" for p in reader.pages)
        except Exception:
            return ""
    if file_type == "docx":
        try:
            import docx
            doc = docx.Document(path)
            return "\n".join(p.text for p in doc.paragraphs)
        except Exception:
            return ""
    if file_type == "xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines)
        except Exception:
            return ""
    return ""


def _chunk_text(text: str, size: int, overlap: int) -> List[str]:
    # Chunk by line, not by a flat word list — joining an entire document's
    # words with single spaces destroys row/line structure, which silently
    # corrupts tabular data (CSV/XLSX rows bleed into each other).
    lines = text.split("\n")
    chunks: List[str] = []
    current: List[str] = []
    count = 0
    for line in lines:
        line_words = len(line.split())
        if count + line_words > size and current:
            chunks.append("\n".join(current))
            tail: List[str] = []
            tail_count = 0
            for prev_line in reversed(current):
                tail_count += len(prev_line.split())
                tail.insert(0, prev_line)
                if tail_count >= overlap:
                    break
            current, count = tail, tail_count
        current.append(line)
        count += line_words
    if current:
        chunks.append("\n".join(current))
    return chunks


async def _fetch_url_text(url: str) -> str:
    from bs4 import BeautifulSoup

    # fetch_with_redirects already re-validates every redirect hop with
    # assert_public_url (see web_crawl.py) -- a business-controlled public
    # URL could otherwise 302 to a private/internal address or cloud
    # metadata endpoint and have that response ingested.
    resp, _chain = await web_crawl.fetch_with_redirects(url, max_bytes=MAX_URL_FETCH_BYTES)
    if resp.status_code >= 400:
        raise HTTPException(status_code=400, detail=f"URL returned status {resp.status_code}")

    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        tag.decompose()
    lines = [line.strip() for line in soup.get_text(separator="\n").splitlines() if line.strip()]
    return "\n".join(lines)


async def ingest_url(db: Session, business_id: str, user_id: str, url: str) -> Document:
    text = await _fetch_url_text(url)

    doc = Document(
        business_id=business_id,
        filename=url,
        file_url=url,
        file_type="url",
        status="processing",
        uploaded_by=user_id,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    _embed_and_store(db, doc, text)
    return doc


async def crawl_and_ingest_website(business_id: str, user_id: str, urls: List[str]) -> None:
    """Background-task worker for POST /api/documents/from-website. Runs
    after the response is already sent, so it opens its own DB session --
    the request's injected session is closed by then (see get_db's
    `finally: db.close()`). Each page reuses ingest_url unchanged; a single
    bad page (404, timeout, odd encoding) is skipped rather than aborting
    the rest of the batch."""
    db = SessionLocal()
    try:
        from ..models.business import Business

        business = db.query(Business).filter(Business.id == business_id).first()
        if not business:
            return

        for url in urls:
            plan = plan_service.get_plan(business.plan)
            if plan.limits.max_document_uploads is not None:
                current_count = db.query(Document).filter(Document.business_id == business_id).count()
                if current_count >= plan.limits.max_document_uploads:
                    break  # plan cap reached mid-crawl -- stop, don't raise (nothing to return this to)

            already_exists = (
                db.query(Document)
                .filter(Document.business_id == business_id, Document.filename == url)
                .first()
            )
            if already_exists:
                continue

            try:
                await ingest_url(db, business_id, user_id, url)
            except Exception:
                continue  # one bad page shouldn't abort the rest of the crawl

            await asyncio.sleep(0.5)  # politeness toward the target site, not a security control
    finally:
        db.close()


async def ingest_document(
    db: Session, business_id: str, user_id: str, file: UploadFile
) -> Document:
    ext = (file.filename or "").rsplit(".", 1)[-1].lower()
    if ext not in ALLOWED_TYPES:
        raise HTTPException(status_code=400, detail=f"File type .{ext} not supported")

    existing = (
        db.query(Document)
        .filter(Document.business_id == business_id, Document.filename == file.filename)
        .first()
    )
    if existing:
        raise HTTPException(
            status_code=400,
            detail=f"'{file.filename}' has already been uploaded. Delete the existing copy first if you want to replace it.",
        )

    content = await file.read()
    if len(content) > settings.max_upload_size_mb * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File too large")

    os.makedirs(settings.upload_dir, exist_ok=True)
    saved_name = f"{uuid.uuid4()}.{ext}"
    saved_path = os.path.join(settings.upload_dir, saved_name)
    with open(saved_path, "wb") as f:
        f.write(content)

    doc = Document(
        business_id=business_id,
        filename=file.filename,
        file_url=saved_path,
        file_type=ext,
        status="processing",
        uploaded_by=user_id,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    _process_document(db, doc, saved_path, ext)
    return doc


def _process_document(db: Session, doc: Document, path: str, ext: str):
    text = _extract_text(path, ext)
    _embed_and_store(db, doc, text)


def _embed_and_store(db: Session, doc: Document, text: str):
    try:
        chunks = _chunk_text(text, settings.chunk_size, settings.chunk_overlap)
        embeddings = embed_texts(chunks)

        for idx, (chunk_text, emb) in enumerate(zip(chunks, embeddings)):
            chunk = DocumentChunk(
                business_id=doc.business_id,
                document_id=doc.id,
                chunk_index=idx,
                content=chunk_text,
                embedding_json=json.dumps(emb),
            )
            db.add(chunk)

        doc.status = "embedded"
    except Exception:
        doc.status = "failed"
    finally:
        db.commit()
